import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)

def set_callout_box(cell, bg_color="F0F4F8", border_color="0F4C81"):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)

def create_report():
    doc = docx.Document()
    
    # Page setup: A4, standard Vietnamese administrative margins
    sections = doc.sections
    for section in sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("PTSC Quảng Ngãi | Báo cáo Chuyển đổi số & Nền tảng Dữ liệu")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Trang ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(9)
        frun.font.color.rgb = RGBColor(120, 120, 120)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 30, 30)

    # --- PHẦN TIÊU NGỮ / HEADER CÔNG TY ---
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_hdr.autofit = False
    
    # Cell 1: Tên đơn vị
    c1 = tbl_hdr.cell(0, 0)
    c1.width = Cm(8.0)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.line_spacing = 1.1
    p1.paragraph_format.space_after = Pt(2)
    r = p1.add_run("TỔNG CÔNG TY CỔ PHẦN DVKT DẦU KHÍ VIỆT NAM\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(9.5)
    r.font.bold = False
    r = p1.add_run("CÔNG TY CỔ PHẦN DỊCH VỤ DẦU KHÍ QUẢNG NGÃI\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r = p1.add_run("TỔ CÔNG TÁC CHUYỂN ĐỔI SỐ")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 76, 129)
    
    # Cell 2: Quốc hiệu / Ngày tháng
    c2 = tbl_hdr.cell(0, 1)
    c2.width = Cm(8.0)
    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing = 1.1
    p2.paragraph_format.space_after = Pt(2)
    r = p2.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r = p2.add_run("Độc lập – Tự do – Hạnh phúc\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r = p2.add_run("Quảng Ngãi, ngày 09 tháng 09 năm 2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.italic = True
    
    # Add small spacing
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(8)
    p_sep.paragraph_format.space_after = Pt(8)

    # --- TIÊU ĐỀ BÁO CÁO ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("BÁO CÁO CHUYÊN ĐỀ")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(15)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 76, 129)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    p_sub.paragraph_format.line_spacing = 1.2
    r_sub = p_sub.add_run("V/v: Kết quả nghiên cứu tài liệu Nền tảng Dữ liệu Tổng công ty\nvà Đề xuất phương án triển khai phần mềm Quản lý Mua sắm, HSEQ tại PTSC Quảng Ngãi")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(30, 41, 59)

    # Kính gửi
    p_kg = doc.add_paragraph()
    p_kg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_kg.paragraph_format.space_after = Pt(14)
    r_kg = p_kg.add_run("Kính gửi: ")
    r_kg.font.bold = True
    r_kg.font.italic = True
    r_kg2 = p_kg.add_run("Ban Giám đốc Công ty Cổ phần Dịch vụ Dầu khí Quảng Ngãi")
    r_kg2.font.bold = True

    # Helper function for section headings
    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 76, 129) # Navy
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_body(text, indent=True, space_after=4, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(space_after)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Cm(0.8)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11.5)
        return p

    def add_callout(text, title=None):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Cm(16.0)
        set_callout_box(cell, bg_color="F1F5F9", border_color="0F4C81")
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if title:
            rt = p.add_run(title + "\n")
            rt.font.bold = True
            rt.font.size = Pt(11)
            rt.font.color.rgb = RGBColor(15, 76, 129)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.italic = True
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ==================== NỘI DUNG BÁO CÁO ====================

    # Căn cứ pháp lý & bối cảnh
    add_callout(
        "Căn cứ chỉ đạo của Ban Giám đốc PTSC Quảng Ngãi về việc rà soát phương án triển khai phần mềm Quản lý Mua sắm hàng hóa dịch vụ và phần mềm HSEQ;\n"
        "Căn cứ ý kiến của Ban Dự án Chuyển đổi số Tổng công ty (email ngày 07/09/2026 của bà Phan Thị Ngọc Vân) đề nghị PTSC Quảng Ngãi nghiên cứu kỹ kiến trúc tổng quan và đề xuất kỹ thuật trong tài liệu Data Platform đã công bố;\n"
        "Tổ công tác Chuyển đổi số PTSC Quảng Ngãi kính báo cáo Ban Giám đốc kết quả nghiên cứu chi tiết và đề xuất kế hoạch hành động cụ thể như sau:",
        "CĂN CỨ VÀ BỐI CẢNH BÁO CÁO"
    )

    # ----------------------------------------------------
    # PHẦN 1
    # ----------------------------------------------------
    add_heading_1("I. BÁO CÁO KẾT QUẢ NGHIÊN CỨU TOÀN DIỆN TÀI LIỆU CỦA TỔNG CÔNG TY")
    add_body(
        "Thực hiện chỉ đạo của Ban Giám đốc và phản hồi từ Ban Dự án Chuyển đổi số (CĐS) Tổng công ty, Tổ công tác đã tiến hành rà soát chuyên sâu 04 bộ tài liệu chính thức do Tổng công ty và Liên danh tư vấn HIPT - AITS công bố, bao gồm: "
        "(1) Báo cáo Đề xuất Kỹ thuật Nền tảng Dữ liệu PTSC (phiên bản 83 trang); "
        "(2) Slide Hội thảo Nền tảng Dữ liệu – Phiên Sáng (Khung kiến trúc & Vai trò); "
        "(3) Slide Hội thảo Nền tảng Dữ liệu – Phiên Chiều (Chính sách, Phân nhóm đơn vị & Lộ trình); "
        "(4) Quy chế Quản trị Dữ liệu Petrovietnam/PTSC (Mô hình 5 cấp & chuẩn PPDM). "
        "Kết quả nghiên cứu cụ thể như sau:"
    )

    add_heading_2("1. Vị thế và phân nhóm triển khai của PTSC Quảng Ngãi")
    add_bullet(
        "Theo Slide 56 Phiên Chiều (Phân nhóm đơn vị & Level triển khai), Tổng công ty phân bổ 17 đơn vị thành viên vào 4 nhóm. Trong đó, PTSC Quảng Ngãi được xếp vào NHÓM 4 – ĐƠN VỊ LỚN (cùng với PTSC M&C, POS, PPS, Thanh Hóa) áp dụng mức kiến trúc L3–L4.",
        "Xếp loại Đơn vị: "
    )
    add_bullet(
        "TCT xác định các đơn vị Nhóm 4 có quy mô sản xuất kinh doanh lớn, sở hữu hạ tầng ứng dụng và cơ sở dữ liệu riêng (như phân xưởng cơ khí Dung Quất, Cảng tổng hợp Dung Quất, đội tàu lai dắt). Do đó, đơn vị có nhu cầu phân tích nội bộ chuyên sâu độc lập, không áp dụng mô hình dùng chung hoàn toàn như các chi nhánh nhỏ nhóm L1.",
        "Đặc thù nghiệp vụ: "
    )

    add_heading_2("2. Mô hình kiến trúc tổng thể Tổng công ty định hướng")
    add_bullet(
        "Tổng công ty áp dụng mô hình Hub - Spoke kết hợp Hybrid (On-premise kết hợp Microsoft Fabric Cloud). Tổng công ty là Hub trung tâm, các đơn vị thành viên Nhóm 4 là Spoke vệ tinh. Dữ liệu giữa các hệ thống không kết nối trực tiếp điểm-nối-điểm mà bắt buộc đi qua Trục tích hợp ESB (Enterprise Service Bus) và hệ thống Quản trị dữ liệu chủ (MDM) của TCT (Báo cáo Kỹ thuật, Trang 33–38).",
        "Mô hình Hub - Spoke: "
    )
    add_bullet(
        "Để phục vụ phân tích nội bộ, TCT sẽ cấp cho PTSC Quảng Ngãi một phân vùng Tenant L3 riêng trên nền tảng Microsoft Fabric (sử dụng Fabric Workspace Capacity). Toàn bộ dữ liệu vận hành chi tiết của Quảng Ngãi lưu trữ tại đây do Đơn vị quản lý, Tổng công ty và các đơn vị bạn không có quyền xem dữ liệu này (Báo cáo Kỹ thuật, Trang 59 & 78).",
        "Phân vùng Tenant L3: "
    )

    add_heading_2("3. Nguyên tắc bảo vệ chủ quyền dữ liệu và quyền hạn của Đơn vị")
    add_bullet(
        "Tại Trang 24 Slide Phiên Sáng, TCT đưa ra cam kết văn bản: 'Dữ liệu không bị lấy tùy ý. Mỗi đơn vị có Landing Zone riêng; dữ liệu chỉ được đưa vào vùng dùng chung sau khi đơn vị kiểm duyệt và phê duyệt'.",
        "Cam kết vùng đệm kiểm duyệt: "
    )
    add_bullet(
        "Tại Trang 22 Slide Phiên Chiều, Lãnh đạo Đơn vị thành viên được xác lập giữ vai trò CHỦ QUẢN DỮ LIỆU CẤP 3 (Data Owner). Lãnh đạo Đơn vị có toàn quyền: (1) Phê duyệt dữ liệu chia sẻ; (2) Bảo vệ dữ liệu cá nhân theo Nghị định 13/2023/NĐ-CP; và (3) Có quyền DỪNG cung cấp, chia sẻ dữ liệu nếu vi phạm quy định hoặc gây mất an toàn kinh doanh.",
        "Quyền tối cao của Lãnh đạo Đơn vị: "
    )

    # ----------------------------------------------------
    # PHẦN 2
    # ----------------------------------------------------
    add_heading_1("II. NHỮNG ĐIỂM PTSC QUẢNG NGÃI HOÀN TOÀN ĐỒNG THUẬN VỚI TỔNG CÔNG TY")
    add_body(
        "Qua đối chiếu với thực tiễn hệ thống công nghệ thông tin hiện hữu tại đơn vị (phần mềm kế toán FAST, quản trị dự án MESx - PMSx, quản lý nhân sự chấm công VTI, hệ thống điều hành Cảng IRTECH), Tổ công tác nhận thấy định hướng của Tổng công ty là hoàn toàn đúng đắn và kiến nghị Ban Giám đốc thống nhất các điểm sau:"
    )

    add_bullet(
        "Email của bà Phan Thị Ngọc Vân ngày 07/09/2026 cảnh báo về việc tích hợp nội bộ giữa MESx – PMSx – FBO sẽ làm gia tăng số lượng API và rủi ro ATTT khi mở rộng sang FAST, eOffice, VTI là hoàn toàn chính xác. Quảng Ngãi nhất trí DỪNG phương án kết nối trực tiếp Point-to-Point, chuyển dịch sang kiến trúc tích hợp tập trung.",
        "1. Thống nhất dừng kết nối trực tiếp (Point-to-Point): "
    )
    add_bullet(
        "Đơn vị ủng hộ chủ trương đưa phân hệ Quản lý Mua sắm và HSEQ vào danh mục ứng dụng vệ tinh chuẩn kết nối với Trục ESB của TCT. Dữ liệu danh mục Khách hàng, Nhà cung cấp sẽ được đồng bộ theo chuẩn MDM toàn Tập đoàn để tránh trùng lặp.",
        "2. Đưa phần mềm Mua sắm & HSEQ vào hệ sinh thái chuẩn: "
    )
    add_bullet(
        "Đơn vị sẵn sàng tiếp nhận phân vùng Fabric Workspace Capacity L3 do TCT cấp phát để xây dựng các báo cáo điều hành phục vụ riêng Ban Giám đốc và các phòng ban chuyên môn, đảm bảo thống nhất công nghệ với TCT.",
        "3. Khai thác hạ tầng phân tích Microsoft Fabric của TCT: "
    )

    # ----------------------------------------------------
    # PHẦN 3
    # ----------------------------------------------------
    add_heading_1("III. PHÂN TÍCH NHỮNG 'KHOẢNG TRỐNG KỸ THUẬT' VÀ THÁCH THỨC THỰC THI TỪ TÀI LIỆU TCT")
    add_body(
        "Mặc dù tài liệu của TCT đã phác thảo khung lý thuyết rất bài bản, tuy nhiên khi áp dụng vào việc triển khai thực tế hợp đồng phần mềm Quản lý Mua sắm và HSEQ của Quảng Ngãi, Tổ công tác phát hiện 04 KHOẢNG TRỐNG KỸ THUẬT LỚN mà tài liệu TCT chưa làm rõ. Đây là cơ sở cốt lõi để Đơn vị yêu cầu TCT hướng dẫn cụ thể:"
    )

    # Tạo bảng 4 khoảng trống
    tbl_gap = doc.add_table(rows=5, cols=3)
    tbl_gap.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_gap.autofit = False
    set_table_borders(tbl_gap)
    
    # Headers
    headers = ["STT & Khoảng trống", "Nội dung trong tài liệu TCT", "Điểm nghẽn thực tế tại PTSC Quảng Ngãi"]
    col_widths = [Cm(3.8), Cm(5.8), Cm(6.4)]
    for i, h in enumerate(headers):
        cell = tbl_gap.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "0F4C81") # Navy
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    gaps_data = [
        (
            "1. Tiêu chuẩn kỹ thuật API\n(API Specification)",
            "Chỉ nêu nguyên lý chung: hỗ trợ REST/JSON, SOAP, Kafka, bảo mật OAuth2/JWT (Trang 33–34 Báo cáo Kỹ thuật, Trang 43 Phiên Sáng).\nChưa có Data Dictionary hay API Schema cụ thể cho Mua sắm/HSEQ.",
            "Nhà thầu phần mềm Mua sắm của QN không có tài liệu đặc tả chuẩn để lập trình 'đón đầu'. Nếu nhà thầu tự viết API bây giờ, sau này TCT ban hành chuẩn khác thì bắt buộc phải đập đi xây lại, gây lãng phí chi phí rất lớn."
        ),
        (
            "2. Mốc thời gian bàn giao Trục ESB\n(Milestone Timeline)",
            "Lộ trình chia thành nhiều giai đoạn (Slide 51 Phiên Chiều). Giai đoạn 1 tập trung xây Hub TCT; Giai đoạn 2 mới khảo sát mở rộng Spoke cho các ĐVTV Nhóm 4.",
            "Nếu Quảng Ngãi 'giãn tiến độ' phần mềm Mua sắm, thì chính xác tháng mấy năm 2026 hoặc quý mấy năm 2027 TCT mới dựng xong ESB và cấp môi trường kết nối thử nghiệm (UAT)? Chưa có mốc thời gian chốt."
        ),
        (
            "3. Kế hoạch chuyển tiếp nghiệp vụ\n(Transition Plan)",
            "Tài liệu tập trung vào trạng thái đích (To-Be Architecture), hoàn toàn chưa đề cập giải pháp chuyển tiếp cho các dự án đang triển khai dở dang tại đơn vị.",
            "Nghiệp vụ mua sắm vật tư công trình Dung Quất và kiểm soát an toàn HSEQ tại Cảng đang chạy hàng ngày. Đơn vị không thể dừng toàn bộ dự án chờ TCT suốt 6–12 tháng mà để cán bộ làm giấy tờ thủ công."
        ),
        (
            "4. Phân định chi phí & Bản quyền\n(Cost & Licensing)",
            "TCT nêu rõ TCT chịu chi phí nền tảng Hub và Core. Đơn vị tự trang trải hạ tầng ảo hóa tại chỗ và chi phí kết nối hệ thống nguồn (Trang 65–70 Báo cáo Kỹ thuật).",
            "Chi phí mua sắm bản quyền phần mềm Mua sắm và chi phí thuê nhà thầu hiệu chỉnh API để nối vào ESB của TCT có được TCT hỗ trợ một phần không, hay Đơn vị phải tự gánh 100% trong dự toán nội bộ?"
        )
    ]

    for row_idx, data in enumerate(gaps_data, start=1):
        bg = "FFFFFF" if row_idx % 2 != 0 else "F8FAFC"
        for col_idx, text in enumerate(data):
            cell = tbl_gap.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(text)
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 76, 129)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                r = p.add_run(text)
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(6)

    # ----------------------------------------------------
    # PHẦN 4
    # ----------------------------------------------------
    add_heading_1("IV. ĐỀ XUẤT PHƯƠNG ÁN TRIỂN KHAI CỦA PTSC QUẢNG NGÃI (PHƯƠNG ÁN 2 PHA)")
    add_body(
        "Để vừa tuân thủ tuyệt đối kiến trúc tích hợp tập trung của Tổng công ty, vừa bảo đảm tiến độ sản xuất kinh doanh tại Nhà máy cơ khí Dung Quất và Cảng biển, Tổ công tác kính đề xuất Ban Giám đốc phương án triển khai phần mềm Quản lý Mua sắm theo 'Chiến lược 2 Pha linh hoạt':"
    )

    add_bullet(
        "Yêu cầu nhà thầu phần mềm tiếp tục hoàn thiện giao diện, luồng phê duyệt mua sắm, quản lý kho bãi và quy trình đánh giá nhà cung cấp nội bộ phục vụ người dùng Quảng Ngãi. Đưa phần mềm vào vận hành độc lập (Stand-alone) để số hóa ngay các khâu nghiệp vụ đang làm giấy tờ thủ công, bảo đảm tiến độ dự án đã cam kết trong hợp đồng.",
        "PHA 1 — Hoàn thiện Quy trình Nghiệp vụ Nội bộ (T09/2026 – T12/2026): "
    )
    add_bullet(
        "Tạm dừng việc lập trình các kết nối trực tiếp rời rạc (point-to-point) giữa phần mềm Mua sắm với FAST hay MESx. Thay vào đó, yêu cầu nhà thầu xây dựng sẵn các Cổng giao tiếp theo chuẩn Module mở (OpenAPI Ready) – sẵn sàng tiếp nhận các thông số kết nối của Trục ESB ngay khi TCT ban hành.",
        "Đóng băng kết nối điểm-nối-điểm: "
    )
    add_bullet(
        "Ngay khi TCT dựng xong Trục ESB và ban hành chuẩn kết nối cho các đơn vị Nhóm 4, Quảng Ngãi sẽ phối hợp với TCT và nhà thầu thực hiện đấu nối phần mềm Mua sắm vào Trục ESB và hệ thống MDM của TCT. Số liệu báo cáo mua sắm sẽ chảy qua Landing Zone để Ban Giám đốc duyệt trước khi đồng bộ lên Hub TCT.",
        "PHA 2 — Đấu nối Trục tích hợp ESB & Data Platform TCT (Q1/2027 – Q2/2027): "
    )

    add_callout(
        "Lợi ích của Phương án 2 Pha:\n"
        "• Đối với Nội bộ: Không bị đứt gãy công tác mua sắm, cán bộ nghiệp vụ có công cụ làm việc ngay, không vi phạm tiến độ hợp đồng đã ký với đối tác.\n"
        "• Đối với Tổng công ty: Thể hiện sự đồng thuận 100% với kiến trúc tập trung, không phát sinh kết nối rác, sẵn sàng hòa mạng ngay khi Hub của TCT vận hành.",
        "ĐÁNH GIÁ HIỆU QUẢ PHƯƠNG ÁN ĐỀ XUẤT"
    )

    # ----------------------------------------------------
    # PHẦN 5
    # ----------------------------------------------------
    add_heading_1("V. KIẾN NGHỊ BAN GIÁM ĐỐC PHÊ DUYỆT VÀ KẾ HOẠCH LÀM VIỆC VỚI TỔNG CÔNG TY")
    add_body(
        "Từ những phân tích chuyên môn và cơ sở pháp lý, kỹ thuật nêu trên, Tổ công tác kính đề xuất Ban Giám đốc phê duyệt các nội dung sau:"
    )

    add_bullet(
        "Phê duyệt chủ trương triển khai phần mềm Quản lý Mua sắm theo Phương án 2 Pha nêu tại Mục IV; giao Tổ CĐS giám sát nhà thầu thiết kế các module mở chờ kết nối Trục ESB.",
        "1. Phê duyệt định hướng triển khai: "
    )
    add_bullet(
        "Chấp thuận nội dung dự thảo công văn / thư điện tử phản hồi chính thức cho Ban Dự án CĐS Tổng công ty (người nhận: bà Phan Thị Ngọc Vân) theo nội dung tại Phụ lục đính kèm.",
        "2. Phê duyệt nội dung phản hồi TCT: "
    )
    add_bullet(
        "Đề nghị TCT bố trí buổi làm việc kỹ thuật 3 bên (BDA CĐS TCT + Liên danh tư vấn HIPT - AITS + PTSC Quảng Ngãi) trong tháng 09/2026 để thống nhất phương án đấu nối kỹ thuật cho phần mềm Mua sắm và HSEQ.",
        "3. Tổ chức cuộc họp kỹ thuật 3 bên: "
    )

    # Chữ ký
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    tbl_sign = doc.add_table(rows=1, cols=2)
    tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sign.autofit = False
    
    cs1 = tbl_sign.cell(0, 0)
    cs1.width = Cm(8.0)
    ps1 = cs1.paragraphs[0]
    ps1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ps1.add_run("PHÊ DUYỆT CỦA BAN GIÁM ĐỐC\n")
    r.font.bold = True
    r.font.size = Pt(11)
    r = ps1.add_run("(Ký và ghi rõ họ tên)")
    r.font.italic = True
    r.font.size = Pt(10)
    
    cs2 = tbl_sign.cell(0, 1)
    cs2.width = Cm(8.0)
    ps2 = cs2.paragraphs[0]
    ps2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ps2.add_run("TỔ CÔNG TÁC CHUYỂN ĐỔI SỐ\n")
    r.font.bold = True
    r.font.size = Pt(11)
    r = ps2.add_run("Tổ trưởng / Product Owner\n\n\n\n")
    r.font.italic = True
    r.font.size = Pt(10)

    # ----------------------------------------------------
    # PHỤ LỤC: DỰ THẢO EMAIL GỬI BÀ VÂN
    # ----------------------------------------------------
    doc.add_page_break()
    
    p_pl = doc.add_paragraph()
    p_pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pl.paragraph_format.space_after = Pt(4)
    r_pl = p_pl.add_run("PHỤ LỤC: DỰ THẢO THƯ ĐIỆN TỬ PHẢN HỒI BAN DỰ ÁN CĐS TỔNG CÔNG TY")
    r_pl.font.bold = True
    r_pl.font.size = Pt(13)
    r_pl.font.color.rgb = RGBColor(15, 76, 129)
    
    p_pl_sub = doc.add_paragraph()
    p_pl_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pl_sub.paragraph_format.space_after = Pt(14)
    r = p_pl_sub.add_run("(Dự thảo gửi bà Phan Thị Ngọc Vân – Ban Dự án Chuyển đổi số PTSC)")
    r.font.italic = True
    r.font.size = Pt(11)

    # Bảng Header Email
    tbl_mail = doc.add_table(rows=3, cols=2)
    tbl_mail.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mail.autofit = False
    set_table_borders(tbl_mail, color="CBD5E1", sz="6")
    
    mail_meta = [
        ("Người nhận (To):", "Chị Phan Thị Ngọc Vân – Ban Dự án Chuyển đổi số Tổng công ty (vanptn@ptsc.com.vn)"),
        ("Đồng kính gửi (Cc):", "Ban Giám đốc PTSC Quảng Ngãi; Liên danh tư vấn HIPT - AITS; Tổ CĐS PTSC QN"),
        ("Tiêu đề (Subject):", "Re: V/v PTSC Quảng Ngãi triển khai phần mềm Quản lý Mua sắm & Kế hoạch tích hợp Data Platform")
    ]
    for idx, (label, val) in enumerate(mail_meta):
        c_lbl = tbl_mail.cell(idx, 0)
        c_lbl.width = Cm(3.8)
        set_cell_background(c_lbl, "F1F5F9")
        set_cell_margins(c_lbl, top=80, bottom=80, left=100, right=100)
        p = c_lbl.paragraphs[0]
        r = p.add_run(label)
        r.font.bold = True
        r.font.size = Pt(10.5)
        
        c_val = tbl_mail.cell(idx, 1)
        c_val.width = Cm(12.2)
        set_cell_background(c_val, "FFFFFF")
        set_cell_margins(c_val, top=80, bottom=80, left=100, right=100)
        p = c_val.paragraphs[0]
        r = p.add_run(val)
        r.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Nội dung email mẫu
    email_body_paragraphs = [
        ("Dear chị Vân và Ban Dự án Chuyển đổi số Tổng công ty,", False),
        ("Cảm ơn ý kiến góp ý rất kịp thời của chị và Ban DA CĐS TCT tại email ngày 07/09/2026 về phương án tích hợp phần mềm tại PTSC Quảng Ngãi.", False),
        ("Thực hiện khuyến nghị của chị, Tổ công tác Chuyển đổi số PTSC Quảng Ngãi cùng các phòng ban nghiệp vụ đã tiến hành nghiên cứu rất kỹ lưỡng toàn bộ tài liệu Nền tảng Dữ liệu (Data Platform) và Báo cáo Đề xuất Kỹ thuật do Tổng công ty và Liên danh tư vấn HIPT - AITS công bố (đặc biệt các nội dung tại Trang 16–38 Báo cáo Kỹ thuật, Slide 23–25 Phiên Sáng và Slide 56 Phiên Chiều).", False),
        ("PTSC Quảng Ngãi hoàn toàn nhất trí với định hướng của TCT về việc tránh kết nối trực tiếp điểm-nối-điểm (point-to-point) giữa các phần mềm nội bộ nhằm giảm thiểu rủi ro an toàn thông tin và bùng nổ API. Đơn vị cam kết sẽ tuân thủ mô hình Hub - Spoke, sẵn sàng kết nối vào Trục tích hợp ESB và tiếp nhận phân vùng Tenant L3 trên Microsoft Fabric của TCT.", False),
        ("Tuy nhiên, để đảm bảo công tác sản xuất kinh doanh tại Nhà máy cơ khí Dung Quất và Cảng biển không bị gián đoạn, đồng thời bảo đảm tính pháp lý của hợp đồng đang triển khai với đối tác phần mềm, PTSC Quảng Ngãi kính đề nghị chị Vân và Ban Dự án CĐS TCT hỗ trợ làm rõ 04 nội dung kỹ thuật cốt lõi sau:", False),
    ]

    for text, bold in email_body_paragraphs:
        add_body(text, indent=False, space_after=5)

    # 4 Câu hỏi chốt hạ
    mail_questions = [
        ("1. Về Tiêu chuẩn kỹ thuật API:", "Hiện tại tài liệu TCT (Trang 43 Phiên Sáng) mới nêu khung tiêu chuẩn chung (REST/JSON/OAuth2). Xin hỏi TCT và Liên danh tư vấn dự kiến khi nào sẽ ban hành Bộ tài liệu đặc tả API chuẩn (OpenAPI/Swagger) và Từ điển dữ liệu chuẩn (Data Dictionary) cho phân hệ Quản lý Mua sắm và HSEQ để đơn vị yêu cầu nhà thầu lập trình đón đầu?"),
        ("2. Về Mốc thời gian kết nối Trục ESB:", "Theo lộ trình phân nhóm (Slide 56 Phiên Chiều), PTSC Quảng Ngãi thuộc Nhóm 4 (Đơn vị lớn L3-L4). Xin TCT cho biết mốc thời gian dự kiến (quý nào, năm nào) Trục ESB của TCT sẽ sẵn sàng môi trường thử nghiệm (UAT/Staging) để Quảng Ngãi thực hiện đấu nối?"),
        ("3. Về Phương án chuyển tiếp vận hành nghiệp vụ:", "Trong thời gian chờ TCT hoàn thiện Trục ESB, PTSC Quảng Ngãi dự kiến cho nhà thầu hoàn thiện quy trình nghiệp vụ nội bộ (Stand-alone) để đưa vào sử dụng trước nhằm giải quyết nhu cầu số hóa cấp bách, đồng thời đóng gói sẵn các module mở (OpenAPI Ready) chờ TCT. Kính xin ý kiến chấp thuận của TCT về phương án này."),
        ("4. Về Buổi làm việc kỹ thuật 3 bên:", "Để giải quyết dứt điểm các vướng mắc kỹ thuật nêu trên, PTSC Quảng Ngãi trân trọng đề xuất TCT bố trí một buổi làm việc kỹ thuật trực tuyến (qua Microsoft Teams) giữa BDA CĐS TCT, Liên danh tư vấn HIPT - AITS và Tổ CĐS PTSC Quảng Ngãi trong khoảng thời gian từ ngày 15/09 đến 18/09/2026.")
    ]

    for q_title, q_desc in mail_questions:
        add_bullet(q_desc, bold_prefix=q_title + " ")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_body("Rất mong sớm nhận được phản hồi và hướng dẫn của chị Vân và Ban Dự án CĐS Tổng công ty để PTSC Quảng Ngãi có cơ sở triển khai các bước tiếp theo.", indent=False)
    add_body("Trân trọng cảm ơn chị,\n\nTỔ CÔNG TÁC CHUYỂN ĐỔI SỐ – PTSC QUẢNG NGÃI", indent=False, space_after=12)

    # Lưu tài liệu
    output_path = r"d:\My Profiles\DataPlatform\bao_cao_nghien_cuu_tai_lieu_tct_va_de_xuat_ke_hoach.docx"
    doc.save(output_path)
    print(f"Successfully saved document to: {output_path}")

if __name__ == "__main__":
    create_report()
